import io

import pytest
from django.contrib.gis.geos import Point
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.geo.models import Community
from apps.imports.models import ImportBatch
from tests.conftest import client_for

pytestmark = pytest.mark.django_db

KEY = "community-locations"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(autouse=True)
def private_uploads(settings, tmp_path):
    from apps.imports.models import ImportBatch
    settings.PRIVATE_UPLOAD_ROOT = str(tmp_path)
    ImportBatch._meta.get_field("file").storage.location = str(tmp_path)


def csv_file(text, name="points.csv"):
    return SimpleUploadedFile(name, text.encode("utf-8"), content_type="text/csv")


def upload(c, f, **extra):
    return c.post("/api/v1/admin/imports/", {"importer": KEY, "file": f, **extra}, format="multipart")


def rows(c, batch_id, status=None):
    params = {"status": status} if status else {}
    return c.get(f"/api/v1/admin/imports/{batch_id}/rows/", params).json()["results"]


# ---------- templates ----------
def test_excel_template_lists_all_24_communities_with_dropdowns(administrator):
    r = client_for(administrator).get(f"/api/v1/admin/importers/{KEY}/template/")
    assert r.status_code == 200 and r["Content-Type"] == XLSX
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Data", "Lists", "Instructions"]
    ws = wb["Data"]
    assert [c.value for c in ws[1]] == ["LGA *", "Community *", "Latitude *", "Longitude *", "Source of location"]
    data = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(data) == 24 and data[0][:2] == ("Brass", "Fish Town")
    assert any("Lists!" in dv.formula1 for dv in ws.data_validations.dataValidation)


def test_csv_template(administrator):
    r = client_for(administrator).get(f"/api/v1/admin/importers/{KEY}/template/", {"file_type": "csv"})
    lines = r.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "LGA *,Community *,Latitude *,Longitude *,Source of location"
    assert len(lines) == 25


def test_importer_list_respects_roles(entry_user, administrator):
    assert client_for(entry_user).get("/api/v1/admin/importers/").json() == []
    keys = [i["key"] for i in client_for(administrator).get("/api/v1/admin/importers/").json()]
    assert keys == [KEY]
    assert client_for(entry_user).get(f"/api/v1/admin/importers/{KEY}/template/").status_code == 404


# ---------- full round trip with Excel ----------
def test_round_trip_excel_upload_preview_commit(administrator):
    c = client_for(administrator)
    wb = load_workbook(io.BytesIO(c.get(f"/api/v1/admin/importers/{KEY}/template/").content))
    ws = wb["Data"]
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "Amassoma":
            row[2].value, row[3].value, row[4].value = 4.97, 6.09, "Field GPS"
    # keep only the filled row so the upload is valid
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, 2).value != "Amassoma":
            ws.delete_rows(r)
    buf = io.BytesIO()
    wb.save(buf)
    resp = upload(c, SimpleUploadedFile("gps.xlsx", buf.getvalue(), content_type=XLSX))
    assert resp.status_code == 201, resp.content
    batch = resp.json()
    assert batch["counts"]["valid"] == 1 and batch["counts"]["total"] == 1

    # nothing changes until commit
    assert Community.objects.get(name="Amassoma").location is None
    done = c.post(f"/api/v1/admin/imports/{batch['id']}/commit/", {}, format="json")
    assert done.status_code == 200, done.content
    assert done.json()["status"] == "committed"
    amassoma = Community.objects.get(name="Amassoma")
    assert (amassoma.location.y, amassoma.location.x) == (4.97, 6.09)
    assert amassoma.location_source == "Field GPS"
    assert amassoma.history.first().history_change_reason.startswith("Bulk upload")
    log = AuditLog.objects.get(action="import")
    assert log.details["imported"] == 1 and log.details["file"] == "gps.xlsx"


# ---------- validation messages ----------
def test_row_errors_are_specific_and_helpful(administrator):
    c = client_for(administrator)
    text = (
        "LGA,Community,Latitude,Longitude\n"
        "Southern Ijaw,Amassoma,4.97,6.09\n"        # 2 ok
        "Yenagoa,Ogbia twn,4.95,6.30\n"             # 3 wrong LGA/typo
        "Ogbia,Ogbia twn,4.95,6.30\n"               # 4 typo -> suggestion
        "Brass,Okpoama,6.3,4.3\n"                   # 5 swapped
        "Nembe,Fantuo,abc,6.5\n"                    # 6 not a number
        "Bayelsa East,Tombia,4.9,6.2\n"             # 7 unknown LGA
        "Sagbama,,5.1,6.0\n"                        # 8 missing community
        "Southern Ijaw,AMASSOMA,4.971,6.091\n"      # 9 duplicate of row 2 (case-insensitive)
    )
    batch = upload(c, csv_file(text)).json()
    assert batch["counts"] == {"valid": 1, "error": 6, "duplicate": 1, "skipped": 0, "imported": 0, "total": 8}
    by_row = {r["row_number"]: r for r in rows(c, batch["id"])}
    assert "is not a community in Yenagoa" in by_row[3]["errors"]["community"][0]
    assert "Did you mean 'Ogbia Town'?" in by_row[4]["errors"]["community"][0]
    assert "swapped" in by_row[5]["errors"]["latitude"][0]
    assert "must be a number" in by_row[6]["errors"]["latitude"][0]
    assert "not a recognised LGA" in by_row[7]["errors"]["lga"][0]
    assert by_row[8]["errors"]["community"] == ["Community is required."]
    assert by_row[9]["status"] == "duplicate" and "row 2" in by_row[9]["errors"]["_row"][0]


def test_correcting_a_row_rechecks_it(administrator):
    c = client_for(administrator)
    batch = upload(c, csv_file("LGA,Community,Latitude,Longitude\nOgbia,Ogbia twn,4.95,6.30\n")).json()
    row = rows(c, batch["id"])[0]
    assert row["status"] == "error"
    fixed = c.patch(f"/api/v1/admin/imports/{batch['id']}/rows/{row['id']}/",
                    {"raw": {"community": "Ogbia Town"}}, format="json").json()
    assert fixed["status"] == "valid" and fixed["corrected"] is True and fixed["errors"] == {}
    assert c.post(f"/api/v1/admin/imports/{batch['id']}/commit/", {}, format="json").status_code == 200
    assert Community.objects.get(name="Ogbia Town").location is not None


def test_commit_refuses_problem_rows_unless_confirmed(administrator):
    c = client_for(administrator)
    text = "LGA,Community,Latitude,Longitude\nSouthern Ijaw,Amassoma,4.97,6.09\nOgbia,Nowhere,4.9,6.3\n"
    batch = upload(c, csv_file(text)).json()
    r = c.post(f"/api/v1/admin/imports/{batch['id']}/commit/", {}, format="json")
    assert r.status_code == 400 and "1 row(s) still have problems" in r.json()["detail"]
    r = c.post(f"/api/v1/admin/imports/{batch['id']}/commit/", {"skip_problem_rows": True}, format="json")
    assert r.status_code == 200
    assert r.json()["counts"] == {"valid": 0, "error": 0, "duplicate": 0, "skipped": 1, "imported": 1, "total": 2}
    # cannot commit twice
    assert c.post(f"/api/v1/admin/imports/{batch['id']}/commit/", {}, format="json").status_code == 400


def test_skip_and_unskip(administrator):
    c = client_for(administrator)
    batch = upload(c, csv_file("LGA,Community,Latitude,Longitude\nOgbia,Nowhere,4.9,6.3\n")).json()
    row = rows(c, batch["id"])[0]
    assert c.post(f"/api/v1/admin/imports/{batch['id']}/rows/{row['id']}/skip/").json()["status"] == "skipped"
    assert c.post(f"/api/v1/admin/imports/{batch['id']}/rows/{row['id']}/unskip/").json()["status"] == "error"


def test_replacing_existing_location_warns(administrator):
    Community.objects.filter(name="Kaiama").update(location=Point(6.30, 5.08, srid=4326))
    c = client_for(administrator)
    batch = upload(c, csv_file("LGA,Community,Latitude,Longitude\nKolokuma/Opokuma,Kaiama,5.10,6.31\n")).json()
    row = rows(c, batch["id"])[0]
    assert row["status"] == "valid"
    assert "Replaces the current location (5.08000, 6.30000)" in row["warnings"]["latitude"][0]


def test_cancel(administrator):
    c = client_for(administrator)
    batch = upload(c, csv_file("LGA,Community,Latitude,Longitude\nOgbia,Ayakoro,4.9,6.3\n")).json()
    assert c.post(f"/api/v1/admin/imports/{batch['id']}/cancel/").json()["status"] == "cancelled"
    assert c.post(f"/api/v1/admin/imports/{batch['id']}/commit/", {}, format="json").status_code == 400
    assert Community.objects.get(name="Ayakoro").location is None


# ---------- file safety ----------
@pytest.mark.parametrize("name,content,message", [
    ("data.xlsx", b"not really excel", "not a real Excel file"),
    ("data.xls", b"\xd0\xcf\x11\xe0old", "Old Excel (.xls)"),
    ("data.pdf", b"%PDF-1.4", "Excel (.xlsx) or CSV"),
    ("data.csv", b"", "empty"),
    ("data.csv", b"LGA,Community,Latitude,Longitude\n", "no data rows"),
    ("data.csv", b"LGA,Community\nOgbia,Ayakoro\n", "required columns are missing: Latitude, Longitude"),
    ("data.csv", b"PK\x03\x04zipzip", "not a plain CSV"),
])
def test_bad_files_rejected_with_clear_message(administrator, name, content, message):
    r = upload(client_for(administrator), SimpleUploadedFile(name, content))
    assert r.status_code == 400
    assert message in r.json()["detail"]
    assert ImportBatch.objects.count() == 0


def test_size_and_row_limits(administrator, settings):
    settings.IMPORT_MAX_ROWS = 2
    text = "LGA,Community,Latitude,Longitude\n" + "Ogbia,Ayakoro,4.9,6.3\n" * 3
    r = upload(client_for(administrator), csv_file(text))
    assert "limit is 2" in r.json()["detail"]


def test_unknown_columns_reported_and_headers_flexible(administrator):
    text = "lga, COMMUNITY ,latitude *,Longitude,Notes\nOgbia,Ayakoro,4.9,6.3,hello\n"
    batch = upload(client_for(administrator), csv_file(text)).json()
    assert batch["counts"]["valid"] == 1
    assert batch["file_warnings"] == ["Column 'Notes' is not part of the template and was ignored."]


def test_semicolon_csv_and_windows_encoding(administrator):
    text = "LGA;Community;Latitude;Longitude;Source of location\nOgbia;Ayakoro;4,9;6,3;Équipe terrain\n"
    # decimal commas are not supported in semicolon files: '4,9' -> 49 is rejected by bounds, not silently accepted
    batch = upload(client_for(administrator),
                   SimpleUploadedFile("x.csv", text.encode("cp1252"), content_type="text/csv")).json()
    row = rows(client_for(administrator), batch["id"])[0]
    assert row["raw"]["location_source"] == "Équipe terrain"
    assert row["status"] == "error"


# ---------- permissions ----------
def test_data_entry_cannot_upload_reference_locations(entry_user):
    r = upload(client_for(entry_user), csv_file("LGA,Community,Latitude,Longitude\nOgbia,Ayakoro,4.9,6.3\n"))
    assert r.status_code == 403


def test_anonymous_blocked(client):
    assert client.get("/api/v1/admin/imports/").status_code in (401, 403)
    assert client.get(f"/api/v1/admin/importers/{KEY}/template/").status_code in (401, 403)

"""Build downloadable upload templates from an importer's column list."""

import csv
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="0B4F6C")
REQUIRED_FILL = PatternFill("solid", fgColor="8A1C1C")


def header_text(col) -> str:
    return f"{col.header} *" if col.required else col.header


def build_xlsx(importer) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    guide = wb.create_sheet("Instructions")

    list_col = 0
    for idx, col in enumerate(importer.columns, start=1):
        letter = get_column_letter(idx)
        cell = ws.cell(row=1, column=idx, value=header_text(col))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if col.required else HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        if col.help:
            cell.comment = Comment(col.help, "Blue Bayelsa")
        ws.column_dimensions[letter].width = col.width
        target = f"{letter}2:{letter}{max(2, 1001)}"

        if col.kind == "choice" and col.choices:
            options = col.choices()
            list_col += 1
            lcol = get_column_letter(list_col)
            for r, opt in enumerate(options, start=1):
                lists.cell(row=r, column=list_col, value=opt)
            dv = DataValidation(type="list", formula1=f"=Lists!${lcol}$1:${lcol}${max(1, len(options))}",
                                allow_blank=not col.required, showErrorMessage=True,
                                errorTitle="Not in the list", error=f"Choose a {col.header} from the list.")
        elif col.kind in ("decimal", "integer") and (col.min_value is not None or col.max_value is not None):
            dv = DataValidation(type="whole" if col.kind == "integer" else "decimal", operator="between",
                                formula1=str(col.min_value if col.min_value is not None else -1e12),
                                formula2=str(col.max_value if col.max_value is not None else 1e12),
                                showErrorMessage=True, errorTitle="Out of range",
                                error=f"{col.header} must be between {col.min_value} and {col.max_value}.")
        elif col.kind == "date":
            dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
                                showErrorMessage=True, error="Enter a date, e.g. 31/03/2026.")
            for r in range(2, 1002):
                ws[f"{letter}{r}"].number_format = "yyyy-mm-dd"
        else:
            dv = None
        if dv is not None:
            ws.add_data_validation(dv)
            dv.add(target)

    keys = [c.key for c in importer.columns]
    for r, row in enumerate(importer.prefill_rows(), start=2):
        for c, key in enumerate(keys, start=1):
            if row.get(key) is not None:
                ws.cell(row=r, column=c, value=row[key])
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32

    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 11
    guide.column_dimensions["C"].width = 70
    guide.column_dimensions["D"].width = 24
    guide["A1"] = importer.label
    guide["A1"].font = Font(bold=True, size=14)
    guide["A2"] = importer.description
    guide["A2"].alignment = Alignment(wrap_text=True)
    guide.merge_cells("A2:D2")
    guide.row_dimensions[2].height = 45
    guide["A3"] = ("Fill in the 'Data' sheet. Columns marked * (red headings) are required. "
                   "Do not rename or reorder the headings. Upload the file on the Bulk upload page; "
                   "you will see every problem before anything is saved.")
    guide["A3"].alignment = Alignment(wrap_text=True)
    guide.merge_cells("A3:D3")
    guide.row_dimensions[3].height = 45
    for c, title in enumerate(["Column", "Required", "What to enter", "Example"], start=1):
        cell = guide.cell(row=5, column=c, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for r, col in enumerate(importer.columns, start=6):
        guide.cell(row=r, column=1, value=col.header)
        guide.cell(row=r, column=2, value="Yes" if col.required else "No")
        guide.cell(row=r, column=3, value=col.help).alignment = Alignment(wrap_text=True)
        guide.cell(row=r, column=4, value=col.example)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(importer) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    keys = [c.key for c in importer.columns]
    writer.writerow([header_text(c) for c in importer.columns])
    for row in importer.prefill_rows():
        writer.writerow(["" if row.get(k) is None else row[k] for k in keys])
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it as UTF-8

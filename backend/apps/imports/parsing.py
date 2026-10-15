"""Safe reading of uploaded .xlsx and .csv files into rows of cell values."""

import csv
import io
import zipfile

from django.conf import settings

XLSX_MAGIC = b"PK\x03\x04"


class FileRejected(Exception):
    """The file cannot be read at all. The message is shown to the user."""


def read_table(uploaded_file) -> tuple[list[str], list[tuple[int, list]]]:
    """Return (headers, [(excel_row_number, cells), ...]) from an upload.

    Checks size, extension and actual file contents (not just the name).
    Excel formulas are never executed: only the values Excel last saved are read.
    """
    name = (uploaded_file.name or "").lower()
    max_bytes = settings.IMPORT_MAX_FILE_MB * 1024 * 1024
    if uploaded_file.size > max_bytes:
        raise FileRejected(f"The file is larger than {settings.IMPORT_MAX_FILE_MB} MB.")
    if uploaded_file.size == 0:
        raise FileRejected("The file is empty.")

    head = uploaded_file.read(4)
    uploaded_file.seek(0)

    if name.endswith(".xlsx"):
        if head != XLSX_MAGIC:
            raise FileRejected("This file is named .xlsx but is not a real Excel file. "
                               "Open it in Excel and use 'Save As' → Excel Workbook (.xlsx).")
        rows = _read_xlsx(uploaded_file)
    elif name.endswith(".csv"):
        if head == XLSX_MAGIC or b"\x00" in uploaded_file.read(4096):
            raise FileRejected("This file is named .csv but is not a plain CSV file.")
        uploaded_file.seek(0)
        rows = _read_csv(uploaded_file.read())
    elif name.endswith(".xls"):
        raise FileRejected("Old Excel (.xls) files are not supported. "
                           "Open the file in Excel and use 'Save As' → Excel Workbook (.xlsx).")
    else:
        raise FileRejected("Please upload an Excel (.xlsx) or CSV (.csv) file.")

    # Drop leading empty rows; the first non-empty row is the header.
    rows = [(n, cells) for n, cells in rows]
    while rows and _is_empty(rows[0][1]):
        rows.pop(0)
    if not rows:
        raise FileRejected("The file has no data.")
    header_row = [("" if c is None else str(c)) for c in rows[0][1]]
    data = [(n, cells) for n, cells in rows[1:] if not _is_empty(cells)]
    if not data:
        raise FileRejected("The file has column headings but no data rows.")
    if len(data) > settings.IMPORT_MAX_ROWS:
        raise FileRejected(f"The file has {len(data)} rows; the limit is {settings.IMPORT_MAX_ROWS} per upload. "
                           "Please split it into smaller files.")
    return header_row, data


def _is_empty(cells) -> bool:
    return all(c is None or str(c).strip() == "" for c in cells)


def _read_xlsx(f):
    from openpyxl import load_workbook

    try:
        wb = load_workbook(f, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise FileRejected("The Excel file could not be opened. It may be damaged or password-protected.")
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
    rows = [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
    wb.close()
    return rows


def _read_csv(data: bytes):
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise FileRejected("The CSV file's text encoding was not recognised. Save it as 'CSV UTF-8'.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [(i, row) for i, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1)]
